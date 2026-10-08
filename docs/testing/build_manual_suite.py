"""Generate SDC Learn Phase 2 manual test Excel suite."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from pathlib import Path

wb = Workbook()

header_fill = PatternFill('solid', fgColor='0F3D6E')
header_font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
should_fill = PatternFill('solid', fgColor='E8F0FE')
pass_fill = PatternFill('solid', fgColor='C6EFCE')
fail_fill = PatternFill('solid', fgColor='FFC7CE')
block_fill = PatternFill('solid', fgColor='FFEB9C')
thin = Border(
    left=Side(style='thin', color='D0D5DD'),
    right=Side(style='thin', color='D0D5DD'),
    top=Side(style='thin', color='D0D5DD'),
    bottom=Side(style='thin', color='D0D5DD'),
)
wrap = Alignment(wrap_text=True, vertical='top')

COLS = [
    ('Case ID', 12), ('Area', 16), ('Feature / Workflow', 28), ('Priority', 10),
    ('Role', 14), ('Account (demo)', 26), ('Preconditions', 34), ('Test steps', 48),
    ('Expected result', 42), ('Actual result', 26), ('Status', 12),
    ('Tester', 12), ('Date', 12), ('Evidence / Notes', 26),
]
STATUS_LIST = '"Pass,Fail,Blocked,N/A,Not Run"'


def add_status_dv(ws, max_row=500):
    dv = DataValidation(type='list', formula1=STATUS_LIST, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'K2:K{max_row}')
    dv2 = DataValidation(type='list', formula1='"Must,Should,Could,NFR"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f'D2:D{max_row}')


def add_rows(ws, rows):
    for i, (title, width) in enumerate(COLS, 1):
        cell = ws.cell(1, i, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'A2'
    for r_i, row in enumerate(rows, 2):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, val)
            cell.alignment = wrap
            cell.border = thin
            cell.font = Font(name='Calibri', size=10)
        if row[3] == 'Must':
            ws.cell(r_i, 4).fill = PatternFill('solid', fgColor='FFF3CD')
        elif row[3] == 'Should':
            ws.cell(r_i, 4).fill = should_fill
        ws.row_dimensions[r_i].height = 68
    ws.auto_filter.ref = f'A1:N{ws.max_row}'
    add_status_dv(ws, max(80, ws.max_row + 20))
    ws.conditional_formatting.add(f'K2:K{ws.max_row}', CellIsRule(operator='equal', formula=['"Pass"'], fill=pass_fill))
    ws.conditional_formatting.add(f'K2:K{ws.max_row}', CellIsRule(operator='equal', formula=['"Fail"'], fill=fail_fill))
    ws.conditional_formatting.add(f'K2:K{ws.max_row}', CellIsRule(operator='equal', formula=['"Blocked"'], fill=block_fill))
    ws.conditional_formatting.add(f'K2:K{ws.max_row}', CellIsRule(operator='equal', formula=['"N/A"'], fill=PatternFill('solid', fgColor='E7E5E4')))


def C(cid, area, feature, pri, role, acct, pre, steps, exp, status='Not Run'):
    return [cid, area, feature, pri, role, acct, pre, steps, exp, '', status, '', '', '']


A, I, L = 'admin@sdclearn.demo', 'instructor@sdclearn.demo', 'learner@sdclearn.demo'
PW = 'Password Demo@123 for all demos.'

# Cover
ws = wb.active
ws.title = '00 Instructions'
ws['A1'] = 'SDC Learn — Manual Test Suite (LMS + Phase 2 SDC Learn AI)'
ws['A1'].font = Font(bold=True, size=16, color='0F3D6E')
ws.merge_cells('A1:F1')
ws['A3'] = 'Purpose'
ws['A3'].font = Font(bold=True, size=12)
ws['A4'] = 'End-to-end manual verification of learner, instructor, coordinator, and SDC Learn AI workflows. One row per case; update Status as you execute.'
ws.merge_cells('A4:F4')
ws['A6'] = 'Environment'
ws['A6'].font = Font(bold=True, size=12)
for i, (k, v) in enumerate([
    ('Base URL', 'http://localhost:3000'),
    ('Login', 'http://localhost:3000/login.html'),
    ('App', 'http://localhost:3000/app.html'),
    ('Cert verify', 'http://localhost:3000/verify.html'),
    ('AI health', 'http://localhost:3000/api/ai/status'),
    ('API health', 'http://localhost:3000/api/health'),
    ('BRD', 'docs/BRD_Phase2_AI_Capabilities.md'),
], 7):
    ws.cell(i, 1, k).font = Font(bold=True)
    ws.cell(i, 2, v)

ws['A16'] = 'Demo accounts (password: Demo@123)'
ws['A16'].font = Font(bold=True, size=12)
ws['A17'] = 'Role'
ws['B17'] = 'Email'
ws['C17'] = 'Use for'
for c in range(1, 4):
    ws.cell(17, c).fill = header_fill
    ws.cell(17, c).font = header_font
for i, row in enumerate([
    ('Coordinator', A, 'Full console, Integrations'),
    ('Instructor', I, 'Excel course, grading, quizzes'),
    ('Learner', L, 'My Courses, tutor, practice, quizzes'),
    ('Accounts Officer', 'accounts@sdclearn.demo', 'Fees-limited custom role'),
], 18):
    for j, v in enumerate(row, 1):
        ws.cell(i, j, v)

ws['A24'] = 'How to use'
ws['A24'].font = Font(bold=True, size=12)
for i, s in enumerate([
    '1. Start server: npm start → confirm /api/health and /api/ai/status.',
    '2. If quizzes/AI seed data missing: Settings → Data → Reset demo data.',
    '3. Configure AI: admin → Settings → Integrations → API key → Test → Save capability map.',
    '4. Execute cases sheet by sheet. Set Status = Pass / Fail / Blocked / N/A / Not Run.',
    '5. Log Actual result + Evidence (screenshot path or note).',
    '6. Open 99 Dashboard for Pass/Fail counts (Excel formulas).',
    '7. Priority Must = Phase 2 release gate; Should = optional; NFR = quality bar.',
], 25):
    ws.cell(i, 1, s)
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)

ws['A33'] = 'Status legend'
ws['A33'].font = Font(bold=True)
ws['A34'] = 'Pass'
ws['A34'].fill = pass_fill
ws['B34'] = 'Fail'
ws['B34'].fill = fail_fill
ws['C34'] = 'Blocked'
ws['C34'].fill = block_fill
ws['D34'] = 'N/A or Not Run'
for col, w in zip('ABCDEF', [22, 36, 40, 18, 18, 40]):
    ws.column_dimensions[col].width = w

sheets = {
    '01 Smoke & Auth': [
        C('SMK-01', 'Smoke', 'Server health', 'Must', 'Any', '—', 'Server running', 'Open /api/health', 'JSON status ok, service SDC Learn API'),
        C('SMK-02', 'Smoke', 'AI API status', 'Must', 'Any', '—', 'Phase 2 deployed', 'Open /api/ai/status', 'ok:true, product SDC Learn AI'),
        C('SMK-03', 'Smoke', 'Login page loads', 'Must', 'Guest', '—', '—', 'Open /login.html', 'SDC branding; demo shortcuts'),
        C('AUTH-01', 'Auth', 'Coordinator login', 'Must', 'Coordinator', A, PW, 'Sign in', 'Staff dashboard'),
        C('AUTH-02', 'Auth', 'Instructor login', 'Must', 'Instructor', I, PW, 'Sign in', 'Instructor dashboard'),
        C('AUTH-03', 'Auth', 'Learner login', 'Must', 'Learner', L, PW, 'Sign in', 'My Courses'),
        C('AUTH-04', 'Auth', 'Wrong password', 'Must', 'Guest', A, '—', 'Wrong password', 'Error; no session'),
        C('AUTH-05', 'Auth', 'Logout', 'Must', 'Learner', L, 'Signed in', 'Sign out', 'Back to login'),
        C('AUTH-06', 'Auth', 'Remember me', 'Should', 'Learner', L, '—', 'Remember me → reopen app', 'Still signed in'),
        C('AUTH-07', 'Auth', 'Theme toggle', 'NFR', 'Any', L, 'Signed in', 'Toggle light/dark', 'Persists'),
        C('AUTH-08', 'Auth', 'Mobile login', 'NFR', 'Guest', '—', '≤375px', 'Open login', 'No overflow; usable'),
    ],
    '02 Learner LMS': [
        C('LRN-01', 'Learner', 'My Courses cards', 'Must', 'Learner', L, 'Enrolled', 'Open My Courses', 'Cards with progress/status'),
        C('LRN-02', 'Learner', 'Course home', 'Must', 'Learner', L, 'Excel AI course', 'Open course', 'Sessions + Outline/Submit/Quizzes pins'),
        C('LRN-03', 'Learner', 'Session video/resources', 'Must', 'Learner', L, 'Unlocked session', 'Open session', 'Video + downloads + prev/next'),
        C('LRN-04', 'Learner', 'Mark complete', 'Must', 'Learner', L, 'Completable', 'Mark complete / undo', 'Progress updates'),
        C('LRN-05', 'Learner', 'Restricted lock', 'Must', 'Learner', 'Restricted enrollment', 'Restricted access', 'Open locked session', 'Locked state'),
        C('LRN-06', 'Learner', 'Zoom panel', 'Must', 'Learner', L, 'Zoom configured', 'Header Zoom', 'URL/ID/passcode + copy'),
        C('LRN-07', 'Learner', 'Help panel', 'Must', 'Learner', L, '—', 'Header Help', 'Support link; separate from Tutor'),
        C('LRN-08', 'Learner', 'Outline', 'Must', 'Learner', L, '—', 'Open Outline', 'Outcomes + roadmap'),
        C('LRN-09', 'Learner', 'Assignment submit', 'Must', 'Learner', L, 'Pending assignment', 'Upload allowed file', 'Confirmed; replace until graded'),
        C('LRN-10', 'Learner', 'Reject bad type', 'Must', 'Learner', L, '—', 'Upload disallowed type', 'Validation error'),
        C('LRN-11', 'Learner', 'Calendar', 'Should', 'Learner', L, 'feature on', 'Open Calendar', 'Sessions listed'),
        C('LRN-12', 'Learner', 'Attendance view', 'Should', 'Learner', L, 'feature on', 'Open Attendance', 'Own records'),
        C('LRN-13', 'Learner', 'Fees', 'Should', 'Learner', L, 'feature on', 'Open Fees', 'Own fees'),
        C('LRN-14', 'Learner', 'Messages', 'Should', 'Learner', L, 'feature on', 'Send/read message', 'Works'),
        C('LRN-15', 'Learner', 'Certificates list', 'Should', 'Learner', L, 'Has CERT-1', 'Open Certificates', 'Valid cert shown'),
        C('LRN-16', 'Learner', 'Public verify', 'Must', 'Guest', '—', 'SDC-2026-0001', 'verify.html?code=…', 'Learner/course/status'),
        C('LRN-17', 'Learner', 'Mid-course feedback', 'Should', 'Learner', L, 'Progress ≥ threshold', 'Submit feedback', 'Saved'),
        C('LRN-18', 'Learner', 'Profile + password', 'Must', 'Learner', L, '—', 'Edit profile; change password', 'Saved; new password works'),
        C('LRN-19', 'Learner', 'Empty states', 'Should', 'Learner', 'No enrollments', '—', 'My Courses', 'Friendly empty state'),
        C('LRN-20', 'Learner', 'Mobile course/session', 'NFR', 'Learner', L, '≤375px', 'Browse course+session', 'Usable'),
    ],
    '03 Instructor LMS': [
        C('INS-01', 'Instructor', 'Dashboard to-grade', 'Must', 'Instructor', I, 'Pending submission', 'Open dashboard', 'Grade list + modal'),
        C('INS-02', 'Instructor', 'Course builder', 'Must', 'Instructor', I, 'Assigned course', 'Manage course', 'Overview stats'),
        C('INS-03', 'Instructor', 'Edit session', 'Must', 'Instructor', I, '—', 'Edit session fields/resources', 'Learner sees update'),
        C('INS-04', 'Instructor', 'Publish course', 'Must', 'Instructor', I, 'can publish', 'Publish/unpublish', 'Status updates'),
        C('INS-05', 'Instructor', 'Assignments CRUD', 'Must', 'Instructor', I, '—', 'Add/edit assignment', 'On learner submit list'),
        C('INS-06', 'Instructor', 'Grade submission', 'Must', 'Instructor', I, 'Submitted file', 'Grade → Save', 'Graded + notify'),
        C('INS-07', 'Instructor', 'No grade without perm', 'Must', 'Accounts', 'accounts@sdclearn.demo', 'No submissions edit', 'Check UI', 'No Grade action'),
        C('INS-08', 'Instructor', 'Attendance', 'Must', 'Instructor', I, 'feature on', 'Mark attendance → save', 'Persisted'),
        C('INS-09', 'Instructor', 'Results entry', 'Must', 'Instructor', I, 'feature on', 'Assessment % → Save/Publish', 'Learner sees result'),
        C('INS-10', 'Instructor', 'Announcement', 'Should', 'Instructor', I, '—', 'Create course announcement', 'Learner sees it'),
        C('INS-11', 'Instructor', 'Course scope', 'Must', 'Instructor', I, 'Assigned only', 'Manage courses list', 'No foreign courses'),
    ],
    '04 Admin Ops': [
        C('ADM-01', 'Admin', 'Learners CRUD', 'Must', 'Coordinator', A, '—', 'Create/edit/deactivate learner', 'Persists'),
        C('ADM-02', 'Admin', 'Enrollment restricted', 'Must', 'Coordinator', A, '—', 'Enroll restricted + allowed sessions', 'Locks correct'),
        C('ADM-03', 'Admin', 'Programs/div/batches', 'Must', 'Coordinator', A, '—', 'CRUD each', 'Lists update'),
        C('ADM-04', 'Admin', 'Roles matrix ai/quizzes', 'Must', 'Coordinator', A, '—', 'Edit role ai + quizzes perms', 'UI respects matrix'),
        C('ADM-05', 'Admin', 'Certificate issue/revoke', 'Must', 'Coordinator', A, 'Eligible', 'Issue then revoke', 'verify.html updates'),
        C('ADM-06', 'Admin', 'Fees', 'Should', 'Coordinator', A, 'feature on', 'Create/update fee', 'Learner sees'),
        C('ADM-07', 'Admin', 'CSV exports', 'Should', 'Coordinator', A, '—', 'Export reports', 'CSV downloads'),
        C('ADM-08', 'Admin', 'Branding settings', 'Must', 'Coordinator', A, '—', 'Change name/colours', 'UI updates'),
        C('ADM-09', 'Admin', 'Terminology', 'Must', 'Coordinator', A, '—', 'Rename Quiz/Learner', 'Labels update'),
        C('ADM-10', 'Admin', 'Feature flags', 'Must', 'Coordinator', A, '—', 'Disable calendar/fees/ai', 'Hidden; LMS core works'),
        C('ADM-11', 'Admin', 'Result weights', 'Must', 'Coordinator', A, '—', 'Set assignment/assessment/quiz/attendance', 'Results text matches'),
        C('ADM-12', 'Admin', 'Backup/restore', 'Must', 'Coordinator', A, '—', 'Download + restore carefully', 'Data restored'),
        C('ADM-13', 'Admin', 'Reset demo', 'Must', 'Coordinator', A, 'Confirm', 'Reset demo data', 'Seed incl. sample quiz'),
        C('ADM-14', 'Admin', 'Backup has no AI keys', 'Must', 'Coordinator', A, 'Key in Integrations', 'Search backup JSON for sk-/AIza', 'No raw keys'),
    ],
    '05 Quizzes Graded': [
        C('QZ-01', 'Quizzes', 'Builder tab', 'Must', 'Instructor', I, 'quizzes feature on', 'Manage → Quizzes', 'Tab lists quizzes'),
        C('QZ-02', 'Quizzes', 'Create manual quiz', 'Must', 'Instructor', I, '—', 'New quiz + MCQ/TF/short → Save', 'Draft saved'),
        C('QZ-03', 'Quizzes', 'Publish', 'Must', 'Instructor', I, 'Draft', 'Publish', 'Learner can see'),
        C('QZ-04', 'Quizzes', 'Unpublish', 'Must', 'Instructor', I, 'Published', 'Unpublish', 'Learner cannot attempt'),
        C('QZ-05', 'Quizzes', 'Learner attempt', 'Must', 'Learner', L, 'QZ-EX-1 published', 'Attempt → submit', 'Objective auto-score'),
        C('QZ-06', 'Quizzes', 'Attempt limit', 'Must', 'Learner', L, 'limit=2', 'Third attempt', 'Blocked'),
        C('QZ-07', 'Quizzes', 'Short answer AI/pending', 'Must', 'Learner', L, 'AI on or off', 'Submit short item', 'AI score or pending note'),
        C('QZ-08', 'Quizzes', 'quizAvg column', 'Must', 'Instructor', I, 'Submitted attempt', 'Results page', 'Quizzes % column'),
        C('QZ-09', 'Quizzes', 'Both assessment+quiz weights', 'Must', 'Coordinator', A, '—', 'Weights + publish result', 'Final uses both'),
        C('QZ-10', 'Quizzes', 'Feature off', 'Must', 'Coordinator', A, '—', 'Disable quizzes', 'Tab/pin hidden'),
        C('QZ-11', 'Quizzes', 'Mobile attempt', 'NFR', 'Learner', L, '≤375px', 'Attempt quiz', 'Sticky submit works'),
        C('QZ-12', 'Quizzes', 'Seed quiz', 'Must', 'Learner', L, 'Fresh demo', 'Excel → Quizzes', 'Module 1 check-in exists'),
    ],
    '06 AI Integrations': [
        C('AI-INT-01', 'Integrations', 'Tab visible', 'Must', 'Coordinator', A, 'settings access', 'Settings → Integrations', 'Provider cards shown'),
        C('AI-INT-02', 'Integrations', 'Master switch', 'Must', 'Coordinator', A, '—', 'Turn AI master off', 'AI UI disabled; LMS ok'),
        C('AI-INT-03', 'Integrations', 'Save OpenAI', 'Must', 'Coordinator', A, 'Valid key', 'Save profile', 'Configured + masked key'),
        C('AI-INT-04', 'Integrations', 'Test OpenAI', 'Must', 'Coordinator', A, 'Saved', 'Test', 'Connected toast'),
        C('AI-INT-05', 'Integrations', 'Save+Test Gemini', 'Must', 'Coordinator', A, 'Valid key', 'Save + Test', 'Success'),
        C('AI-INT-06', 'Integrations', 'Save+Test Azure', 'Must', 'Coordinator', A, 'Endpoint+deployment+key', 'Save + Test', 'Success'),
        C('AI-INT-07', 'Integrations', 'Capability map', 'Must', 'Coordinator', A, '≥1 provider', 'tutor≠evaluate providers; Save', 'Persists reload'),
        C('AI-INT-08', 'Integrations', 'Budget hard stop', 'Must', 'Coordinator', A, 'maxCalls=1 hard stop', '2 AI calls', '2nd blocked; LMS ok'),
        C('AI-INT-09', 'Integrations', 'PII strip on', 'Must', 'Coordinator', A, 'strip enabled', 'Use AI with PII in context', 'No crash; scrub active'),
        C('AI-INT-10', 'Integrations', 'Bad key soft fail', 'Must', 'Coordinator', A, '—', 'Test bad key', 'Error toast; no crash'),
        C('AI-INT-11', 'Integrations', 'No configure perm', 'Must', 'Instructor', I, 'use only', 'Integrations UI', 'View-only / no save'),
        C('AI-INT-12', 'Integrations', 'AI unavailable', 'Must', 'Learner', L, 'No provider', 'Open Tutor', 'Clear error; course works'),
        C('AI-INT-13', 'Integrations', 'Env key fallback', 'Should', 'Coordinator', A, 'Key in .env', 'Status env flags; Test', 'Works without UI key'),
    ],
    '07 AI Tutor & Summary': [
        C('AI-T-01', 'Tutor', 'Header button', 'Must', 'Learner', L, 'aiTutor on', 'Course/session tools', 'SDC Learn AI beside Help/Zoom'),
        C('AI-T-02', 'Tutor', 'Grounded + citations', 'Must', 'Learner', L, 'Provider OK', 'Ask session topic', 'Grounded answer / citations'),
        C('AI-T-03', 'Tutor', 'Out-of-scope refuse', 'Must', 'Learner', L, '—', 'Ask unrelated topic', 'Not-covered refusal'),
        C('AI-T-04', 'Tutor', 'Helpful/Unhelpful', 'Must', 'Learner', L, 'Answer shown', 'Rate response', 'Toast ack'),
        C('AI-T-05', 'Tutor', 'Escalate', 'Must', 'Learner', L, '—', 'Escalate', 'Instructor notification'),
        C('AI-T-06', 'Tutor', 'Disclaimer', 'Must', 'Learner', L, '—', 'Open drawer', 'Not official grading disclaimer'),
        C('AI-T-07', 'Tutor', 'Help coexistence', 'Must', 'Learner', L, '—', 'Use Help and Tutor', 'Both work'),
        C('AI-T-08', 'Tutor', 'No ai:use', 'Must', 'Accounts', 'accounts@sdclearn.demo', '—', 'No tutor / denied', 'Blocked'),
        C('AI-T-09', 'Summarize', 'Summarize resource', 'Must', 'Learner', L, 'aiSummarize on', 'Sparkles on resource', 'English summary'),
        C('AI-T-10', 'Summarize', 'Pin summary', 'Must', 'Instructor', I, 'Manage course', 'Pin summary', 'Visible on session'),
        C('AI-T-11', 'Tutor', 'Mobile drawer', 'NFR', 'Learner', L, '≤375px', 'Chat in drawer', 'Usable full width'),
        C('AI-T-12', 'Tutor', 'English only', 'Must', 'Learner', L, '—', 'Request Urdu', 'English response (P2)'),
    ],
    '08 AI Practice & Generate': [
        C('AI-P-01', 'Practice', 'Generate practice', 'Must', 'Learner', L, 'aiPractice on', 'Test my understanding', 'Private quiz form'),
        C('AI-P-02', 'Practice', 'Not in gradebook', 'Must', 'Learner', L, 'Completed practice', 'Check Results quizAvg', 'Practice not included'),
        C('AI-P-03', 'Practice', 'Explanations', 'Must', 'Learner', L, '—', 'Check answers', 'Score + rationales'),
        C('AI-P-04', 'Practice', 'Flag off', 'Must', 'Coordinator', A, '—', 'Disable aiPractice', 'Button hidden'),
        C('AI-G-01', 'Generator', 'Generate draft', 'Must', 'Instructor', I, 'aiQuizGen + provider', 'Generate with SDC Learn AI', 'Draft items; unpublished'),
        C('AI-G-02', 'Generator', 'Accept selected', 'Must', 'Instructor', I, 'Draft list', 'Uncheck some; accept', 'Only selected saved'),
        C('AI-G-03', 'Generator', 'Never auto-publish', 'Must', 'Instructor', I, 'After accept', 'Check status', 'draft until Publish'),
        C('AI-G-04', 'Checker', 'Short AI score', 'Must', 'Learner', L, 'Short Q + AI', 'Submit quiz', 'aiReviews present'),
        C('AI-G-05', 'Checker', 'Human override', 'Should', 'Instructor', I, 'AI short score', 'Override if UI exists', 'Human final; note gap if missing'),
    ],
    '09 AI Evaluator': [
        C('AI-E-01', 'Evaluator', 'Button in modal', 'Must', 'Instructor', I, 'aiEvaluate on', 'Grade modal', 'Evaluate button'),
        C('AI-E-02', 'Evaluator', 'Draft fills', 'Must', 'Instructor', I, 'Provider + text-like file', 'Evaluate', 'Feedback/marks draft + hint'),
        C('AI-E-03', 'Evaluator', 'Save to publish', 'Must', 'Instructor', I, 'Draft filled', 'Close vs Save', 'Only Save grades/notifies'),
        C('AI-E-04', 'Evaluator', 'aiDraft audit', 'Must', 'Instructor', I, 'Saved with AI', 'Backup JSON submission', 'aiDraft/aiAssisted set'),
        C('AI-E-05', 'Evaluator', 'Binary checklist', 'Must', 'Instructor', I, 'xlsx/pbix file', 'Evaluate', 'Checklist / no invented grade'),
        C('AI-E-06', 'Evaluator', 'No auto-apply', 'Must', 'Coordinator', A, '—', 'Integrations UI', 'No auto-apply grade control'),
        C('AI-E-07', 'Evaluator', 'AI-assisted label', 'Should', 'Learner', L, 'AI-graded work', 'View feedback', 'AI-assisted mention if present'),
    ],
    '10 At-Risk & Results': [
        C('AI-R-01', 'At-Risk', 'Dashboard refresh', 'Must', 'Instructor', I, 'aiAtRisk on', 'Refresh advisories', 'Cards or empty; staff only'),
        C('AI-R-02', 'At-Risk', 'Explainable factors', 'Must', 'Instructor', I, 'Low progress/missing work', 'Refresh', 'Factors list signals'),
        C('AI-R-03', 'At-Risk', 'Hidden from learner', 'Must', 'Learner', L, 'Flags exist', 'Browse learner UI', 'No risk stigma'),
        C('AI-R-04', 'At-Risk', 'Outreach note', 'Must', 'Instructor', I, 'Flag exists', 'Add note', 'Saved on flag'),
        C('AI-R-05', 'At-Risk', 'No auto penalty', 'Must', 'Instructor', I, '—', 'After refresh', 'Grades/enrollment unchanged'),
        C('RES-01', 'Results', 'Quiz in formula text', 'Must', 'Instructor', I, '—', 'Results header', 'Mentions quizzes weight'),
        C('RES-02', 'Results', 'Publish notifies', 'Must', 'Instructor', I, '—', 'Publish result', 'Learner My results updates'),
    ],
    '11 Permissions NFR': [
        C('PERM-01', 'Perms', 'Configure admin-only', 'Must', 'Instructor', I, '—', 'Save Integrations key', 'Denied/view-only'),
        C('PERM-02', 'Perms', 'Learner ai:use', 'Must', 'Learner', L, 'Seed grants', 'Tutor works', 'Success'),
        C('PERM-03', 'Perms', 'Master ai feature off', 'Must', 'All', A, '—', 'features.ai off', 'All AI entry points gone'),
        C('NFR-01', 'NFR', 'Integrations mobile', 'NFR', 'Coordinator', A, '≤375px', 'Integrations', 'Cards stack; usable'),
        C('NFR-02', 'NFR', 'Drawer themes', 'NFR', 'Learner', L, 'Light+dark', 'Tutor drawer', 'Readable'),
        C('NFR-03', 'NFR', 'AI badge', 'NFR', 'Any', L, 'AI output', 'Check badge', 'SDC Learn AI badge'),
        C('NFR-04', 'NFR', 'Keyboard quiz/modal', 'NFR', 'Any', L, '—', 'Tab/submit', 'Focus + submit OK'),
        C('NFR-05', 'NFR', 'LMS during AI wait', 'NFR', 'Learner', L, 'Slow AI', 'Navigate while loading', 'LMS responsive'),
    ],
}

for name, rows in sheets.items():
    w = wb.create_sheet(name)
    add_rows(w, rows)

# Dashboard
dash = wb.create_sheet('99 Dashboard', 1)
dash['A1'] = 'Execution dashboard'
dash['A1'].font = Font(bold=True, size=14, color='0F3D6E')
dash['A2'] = 'Counts use COUNTIF on each case sheet column K. Fill Status values then review totals.'
headers = ['Sheet', 'Total cases', 'Pass', 'Fail', 'Blocked', 'Not Run', 'N/A', 'Pass %']
for i, h in enumerate(headers, 1):
    dash.cell(4, i, h).fill = header_fill
    dash.cell(4, i).font = header_font
names = list(sheets.keys())
for i, name in enumerate(names, 5):
    ref = f"'{name}'!K:K"
    dash.cell(i, 1, name)
    dash.cell(i, 2, f"=COUNTA('{name}'!A:A)-1")
    dash.cell(i, 3, f'=COUNTIF({ref},"Pass")')
    dash.cell(i, 4, f'=COUNTIF({ref},"Fail")')
    dash.cell(i, 5, f'=COUNTIF({ref},"Blocked")')
    dash.cell(i, 6, f'=COUNTIF({ref},"Not Run")')
    dash.cell(i, 7, f'=COUNTIF({ref},"N/A")')
    dash.cell(i, 8, f'=IF(B{i}=0,0,C{i}/B{i})')
    dash.cell(i, 8).number_format = '0.0%'
last = 4 + len(names)
dash.cell(last + 1, 1, 'TOTAL').font = Font(bold=True)
for col in range(2, 8):
    letter = get_column_letter(col)
    cell = dash.cell(last + 1, col, f'=SUM({letter}5:{letter}{last})')
    cell.font = Font(bold=True)
dash.cell(last + 1, 8, f'=IF(B{last+1}=0,0,C{last+1}/B{last+1})')
dash.cell(last + 1, 8).number_format = '0.0%'
dash.cell(last + 1, 8).font = Font(bold=True)
dash[f'A{last+3}'] = 'Release gate: all Must cases in sheets 05–10 Pass (or N/A with reason). Any Fail blocks Phase 2 release.'
dash[f'A{last+5}'] = 'Sign-off'
dash[f'B{last+5}'] = 'Tester:'
dash[f'B{last+6}'] = 'Date:'
dash[f'B{last+7}'] = 'Build/commit:'
dash[f'B{last+8}'] = 'Result (Pass/Fail):'
for col, w in zip('ABCDEFGH', [28, 12, 10, 10, 12, 12, 10, 10]):
    dash.column_dimensions[col].width = w

# Traceability
tr = wb.create_sheet('Traceability')
tr['A1'] = 'BRD / Plan traceability'
tr['A1'].font = Font(bold=True, size=14, color='0F3D6E')
for i, h in enumerate(['BRD / Wave item', 'Test sheet', 'Case ID prefix'], 1):
    tr.cell(3, i, h).fill = header_fill
    tr.cell(3, i).font = header_font
for i, row in enumerate([
    ('W0 Integrations + proxy + secrets + flags', '06 AI Integrations', 'AI-INT'),
    ('W1 Graded quizzes + quizAvg', '05 Quizzes Graded', 'QZ'),
    ('W2 Generator / Checker / Practice', '08 AI Practice & Generate', 'AI-P / AI-G'),
    ('W3 Evaluator draft-only', '09 AI Evaluator', 'AI-E'),
    ('W4 Tutor + English summarizer', '07 AI Tutor & Summary', 'AI-T'),
    ('W5 At-risk advisories', '10 At-Risk & Results', 'AI-R'),
    ('LMS learner delivery', '02 Learner LMS', 'LRN'),
    ('Instructor ops', '03 Instructor LMS', 'INS'),
    ('Admin / settings / backup', '04 Admin Ops', 'ADM'),
    ('Auth / smoke', '01 Smoke & Auth', 'SMK / AUTH'),
    ('Permissions & NFR', '11 Permissions NFR', 'PERM / NFR'),
    ('Deferred: Urdu, AI-08, remedial path', '—', 'Out of scope P2'),
], 4):
    for j, v in enumerate(row, 1):
        cell = tr.cell(i, j, v)
        cell.border = thin
tr.column_dimensions['A'].width = 48
tr.column_dimensions['B'].width = 28
tr.column_dimensions['C'].width = 22

out = Path(__file__).resolve().parent / 'SDC-Learn-Manual-Test-Suite-Phase2.xlsx'
wb.save(out)
print('Wrote', out)
print('Sheets:', wb.sheetnames)
print('Total cases:', sum(len(v) for v in sheets.values()))
