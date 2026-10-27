"""Builds docs/testing/SDC-Learn-Manual-Test-Suite-Full.xlsx: every manual test for the whole system.
Run: python docs/testing/build_full_manual_suite.py   (needs openpyxl)
Edit the CASES lists below when features change; the dashboard counts update by formula."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

OUT = os.path.join(os.path.dirname(__file__), 'SDC-Learn-Manual-Test-Suite-Full.xlsx')
A, I, L, AC, G = 'admin@sdclearn.demo', 'instructor@sdclearn.demo', 'learner@sdclearn.demo', 'accounts@sdclearn.demo', 'Guest (signed out)'
M, S, N = 'Must', 'Should', 'NFR'

# (id, feature, priority, role/account, preconditions, steps, expected)
SHEETS = {
'01 Sign-in & Session': [
 ('AUTH-01', 'Coordinator sign-in', M, A, 'Password Demo@123', '1. Open /login.html  2. Enter email + password  3. Sign in', 'Staff dashboard opens; name shows top right'),
 ('AUTH-02', 'Instructor sign-in', M, I, '', 'Sign in', 'Instructor dashboard with assigned courses only'),
 ('AUTH-03', 'Learner sign-in', M, L, '', 'Sign in', 'My Courses opens'),
 ('AUTH-04', 'Demo account shortcuts', S, G, 'Demo accounts shown', 'Click Coordinator / Instructor / Learner chip, then Sign in', 'Fields fill; sign-in succeeds'),
 ('AUTH-05', 'Wrong password', M, G, '', 'Valid email, wrong password', '"Email or password is incorrect."; no session'),
 ('AUTH-06', 'Unknown email', M, G, '', 'Unregistered email', 'Same generic error (does not reveal which emails exist)'),
 ('AUTH-07', 'Inactive account', M, A, 'Set a learner to Inactive (People)', 'Sign in as that learner', '"This account is not active…"; blocked'),
 ('AUTH-08', 'Remember me', S, L, '', 'Tick Keep me signed in, sign in, close and reopen the browser', 'Still signed in'),
 ('AUTH-09', 'Session survives reload', M, L, 'Signed in', 'Reload on a course page', 'Same page, still signed in'),
 ('AUTH-10', 'Sign out clears data', M, L, 'Signed in', 'User menu → Sign out', 'Back on login; reopening app.html asks to sign in'),
 ('AUTH-11', 'Direct URL when signed out', M, G, 'Signed out', 'Open app.html#/dashboard', 'Redirected to sign-in'),
 ('AUTH-12', 'Forgot password info', S, G, '', 'Click Forgot password?', 'Dialog with coordinator email/phone'),
 ('AUTH-13', 'Continue with Google (registered email)', S, G, 'Google enabled in Supabase', 'Continue with Google → pick an account whose email is registered', 'Signed in to the matching platform account'),
 ('AUTH-14', 'Continue with Google (unknown email)', S, G, 'No Google sign-up role set', 'Use an unregistered Google account', 'Clear "isn\'t registered" message; no account created'),
 ('AUTH-15', 'Privacy and Terms pages', S, G, '', 'Open the links under the sign-in card', 'Pages load with SDC branding'),
],
'02 Branding, UI & Access': [
 ('UI-01', 'SDC branding everywhere', M, G, '', 'Check login, app, verify page, footer', 'SDC Learn / SDC Karachi branding; no "EAD"'),
 ('UI-02', 'Contact details in footer', S, A, '', 'Read the footer', '(021) 99334387, 99334388 · sdckar@sdckarachi.org.pk · website link'),
 ('UI-03', 'Light / dark theme', N, L, '', 'Toggle theme, reload', 'Theme switches with animation and persists'),
 ('UI-04', 'Collapsible sidebar', S, A, 'Desktop', 'Click the logo; collapse a nav group; reload', 'Icon-only sidebar and collapsed groups are remembered'),
 ('UI-05', 'Mobile layout (375 px)', N, L, 'Phone or 375 px window', 'Open login, My Courses, a course, a session, a quiz', 'No sideways scrolling; menu opens as a drawer; buttons usable'),
 ('UI-06', 'Tables paginate', S, A, '20+ learners', 'Open Learners; search, sort, page 2', '"Showing X–Y of Z"; search and sort work'),
 ('UI-07', 'Keyboard: dialogs', N, A, '', 'Open any dialog; press Tab repeatedly; press Escape', 'Focus stays inside; Escape closes only the top dialog and returns focus'),
 ('UI-08', 'Saved confirmation', S, A, '', 'Save any Settings tab', '"Settings saved." toast is visible'),
 ('UI-09', 'Learner cannot open staff pages', M, L, '', 'Type #/dashboard, #/settings, #/roles in the address bar', '"Not available" page'),
 ('UI-10', 'Instructor cannot open Settings', M, I, '', 'Open #/settings', '"Not available" page'),
 ('UI-11', 'Notifications', S, L, 'Has unread notifications', 'Bell icon and Notifications menu item', 'Count badge on both; marking read updates the counts'),
],
'03 Learner Portal': [
 ('LRN-01', 'My Courses', M, L, 'Enrolled', 'Open My Courses; search "Excel"; search "Photoshop"; clear', 'Cards with progress; search filters; empty-state message; clear restores'),
 ('LRN-02', 'Pinned announcements', S, L, 'A pinned announcement exists', 'Open My Courses', 'Banner at the top'),
 ('LRN-03', 'Course home', M, L, '', 'Open a course', 'Sessions with status badges; Outline / Submit / Quizzes pins; instructor of MY batch'),
 ('LRN-04', 'Session page', M, L, '', 'Open a past session', 'Video, resources (Download + Summarize), Previous/Next, Zoom/Help in header'),
 ('LRN-05', 'Mark complete only after the session', M, L, '', 'Open an upcoming session; then a past one', 'Upcoming: note instead of button. Past: Mark as complete updates progress'),
 ('LRN-06', 'Restricted access', M, L, 'Enrollment with selected sessions only', 'Open a session not in the list', 'Locked badge; no video or resources'),
 ('LRN-07', 'Zoom panel', M, L, 'Course has Zoom details', 'Header → Zoom; copy Meeting ID', 'Link, ID, passcode; "copied" toast'),
 ('LRN-08', 'Help panel', S, L, '', 'Header → Help', 'Support form or email link'),
 ('LRN-09', 'Course outline', S, L, '', 'Open Outline; Print', 'Outcomes, modules, delivery; print view'),
 ('LRN-10', 'Submit assignment', M, L, 'Pending assignment in my batch', 'Submit page → choose assignment → type my email → attach .xlsx → Upload', '"Submission received"; instructor notified'),
 ('LRN-11', 'Wrong email / bad file / too big', M, L, '', 'Try a wrong email, a .exe, a 35 MB file', 'Each refused with a clear message'),
 ('LRN-12', 'Replace until graded', M, L, 'Submitted, not graded', 'Upload a new version; then try on a graded one', 'Replaced (history kept); graded work is locked'),
 ('LRN-13', 'Late rules', S, L, 'Past-due assignments: one late allowed, one not', 'Open both', 'Late allowed → flagged Late; not allowed → closed'),
 ('LRN-14', 'Only my batch\'s work', M, L, 'Another batch of the same course has its own assignment', 'Check My Assignments, course home, calendar', 'Other batch\'s work never appears; work added for a new batch does not reach completed learners'),
 ('LRN-15', 'My assignments / calendar / attendance', S, L, '', 'Open each page', 'Own data only; correct statuses'),
 ('LRN-16', 'My results', M, L, 'A published result', 'Open Results', 'Assignments, Quizzes, Assessment, Attendance, Final %, grade'),
 ('LRN-17', 'Certificates', M, L, 'Has a valid certificate', 'View certificate; open its verify link', 'Printable certificate with code; verify page says Valid'),
 ('LRN-18', 'Fees', S, L, '', 'Open Fees', 'Only my fee records; billed / paid / outstanding'),
 ('LRN-19', 'Mid-course feedback', S, L, 'Progress above threshold', 'Give 5 stars + comment', 'Saved; prompt not shown again'),
 ('LRN-20', 'Profile and password', M, L, '', 'Edit phone; change password (wrong current, then right)', 'Wrong current refused; new password works at next sign-in'),
 ('LRN-21', 'Messages', S, L, '', 'Send to instructor; instructor replies', 'Unread dot and count for the receiver'),
],
'04 Instructor': [
 ('INS-01', 'Course scope', M, I, '', 'Open My Courses / Manage courses', 'Only courses where I teach a batch'),
 ('INS-02', 'Course outline (sessions)', M, I, '', 'Course Builder → Sessions: add, edit, reorder, attach a resource file', 'Learners see the change'),
 ('INS-03', 'Batch workspace', M, I, '', 'Course Builder → Batches → Open', 'Assignments, Quizzes, Learners of that batch only'),
 ('INS-04', 'Assignment in a batch', M, I, '', 'Batch → Assignments → Add (session, due, marks)', 'Only that batch\'s learners see it and are notified'),
 ('INS-05', 'Copy from another batch', S, I, 'Another batch has work', 'Copy from another batch', 'Copies arrive as drafts with due dates cleared'),
 ('INS-06', 'Grade a submission', M, I, 'Pending submission', 'Submissions → Grade → marks + feedback → Save', 'Graded; learner notified'),
 ('INS-07', 'Attendance on the day only', M, I, 'A session today and a past one', 'Mark today\'s; open the past one', 'Today saves; past is view-only with a note'),
 ('INS-08', 'Results', M, I, '', 'Results → assessment % → Publish → Save', 'Final uses the configured weights; learner notified'),
 ('INS-09', 'Course announcement', S, I, '', 'New announcement for a course', 'Only enrolled learners see it'),
 ('INS-10', 'At-risk advisories', S, I, '', 'Dashboard → Refresh; add an outreach note', 'Explainable factors; note saved; nothing changes for learners'),
],
'05 Coordinator & Admin': [
 ('ADM-01', 'Programs / divisions / batches', M, A, '', 'Create, edit, delete each', 'Lists update; delete blocked while in use'),
 ('ADM-02', 'Program fills division and end date', S, A, '', 'New course → choose program → set start date', 'Division set; end date from program duration'),
 ('ADM-03', 'Batch names', S, A, '', 'New batch → course + delivery + start date', 'Name auto-generated (e.g. Excel-Onsite-2027)'),
 ('ADM-04', 'Course is a blueprint', M, A, '', 'Open a course in the Course Builder', 'Overview / Sessions / Content / Batches / Learners; no course-level assignments or quizzes'),
 ('ADM-05', 'Enrollment needs a batch', M, A, 'Course has batches', 'Enroll without a batch; then with one; enroll the same learner again', 'Batch required; duplicate refused'),
 ('ADM-06', 'Learners / instructors CRUD', M, A, '', 'Create (password ≥ 8), edit, set Inactive', 'Saved; password stored hashed'),
 ('ADM-07', 'Delete user rules', M, A, '', 'Delete a learner with fees; then one without', 'With fees: refused ("set Inactive"); without: deleted with their records'),
 ('ADM-08', 'Roles: create and edit', M, A, '', 'New role, tick permissions, assign to a user', 'User sees exactly those menus and buttons'),
 ('ADM-09', 'Built-in Coordinator role', M, A, '', 'Edit the Coordinator role', 'Locked, with note to use Duplicate'),
 ('ADM-10', 'Role ceiling', M, AC, 'Accounts Officer can edit roles (grant it first)', 'Try to grant a permission the Accounts Officer does not have', 'Refused: cannot grant what you do not have'),
 ('ADM-11', 'Certificates: eligibility, issue, revoke', M, A, '', 'Issue to an eligible learner; try an ineligible one; revoke with a reason', 'Ineligible disabled; code SDC-YYYY-NNNN; verify shows Revoked'),
 ('ADM-12', 'Fees', S, A, '', 'Create fee; mark paid with reference', 'Pending → Paid with reference'),
 ('ADM-13', 'Announcements to everyone (pinned)', S, A, '', 'New announcement, Everyone, pinned', 'Shown on every dashboard and learner home'),
 ('ADM-14', 'Reports and CSV exports', S, A, '', 'Reports → export each CSV', 'Files download; enrollments/results include reg. no. and marks'),
 ('ADM-15', 'Feedback page', S, A, '', 'Feedback → filter by course and rating', 'Learner feedback with averages; delete works'),
],
'06 Settings': [
 ('SET-01', 'Branding', M, A, '', 'Change product name and colours; save', 'Applied across the app and login'),
 ('SET-02', 'Terminology', M, A, '', 'Rename Learner → Trainee; save; restore defaults', 'Labels change everywhere; defaults restore'),
 ('SET-03', 'Learning rules', M, A, '', 'Change result weights, pass %, grade bands, upload limits', 'Results text and grades follow the new values'),
 ('SET-04', 'Feature switches', M, A, '', 'Flip Calendar off and on', 'Saves instantly with a toast; menu item hides/returns'),
 ('SET-05', 'Sign-in tab', S, A, '', 'Check provider status; copy redirect URL', 'Shows whether Google is enabled; URL copies'),
 ('SET-06', 'Backup and restore', M, A, '', 'Download backup; change something; restore the file', 'Data returns to the backup; backup has no AI keys'),
 ('SET-07', 'Reset demo data', S, A, 'Not on live data', 'Data → Reset to demo data → type RESET → Reset all data', 'Button stays disabled until RESET is typed; demo catalogue (incl. sample quiz) restored'),
 ('SET-08', 'Unsaved changes', S, A, '', 'Edit a field on Learning; click another tab; Cancel; then Save', '"Unsaved changes" shows by the Save button; leaving asks first; Cancel keeps the edit'),
 ('SET-09', 'One AI switch', M, A, '', 'Turn SDC Learn AI off in Features; open Integrations; turn it on there; back to Features', 'Both places show the same state; AI feature switches are greyed out while AI is off'),
 ('SET-10', 'Result weight shares', S, A, '', 'Learning → change a weight', 'The line under the weights shows the real share of each part of the final result'),
 ('SET-11', 'Settings on a phone', N, A, '375 px wide', 'Open each Settings tab, including Integrations', 'Open tab visible in the tab bar; no sideways scrolling; feature map shows Provider and Model labels'),
],
'07 Quizzes': [
 ('QZ-01', 'Build a quiz', M, I, 'Batch workspace', 'New quiz: MCQ (3 options), True/False, Short answer; save', 'Draft saved with the right types'),
 ('QZ-02', 'Validation', S, I, '', 'Save an MCQ with no correct option', 'Refused: mark the correct answer'),
 ('QZ-03', 'Publish / unpublish', M, I, '', 'Publish, then unpublish', 'Learners of the batch can / cannot attempt'),
 ('QZ-04', 'Attempt and score', M, L, 'Published quiz in my batch', 'Answer and submit', 'Objective items scored instantly'),
 ('QZ-05', 'Attempt limit', M, L, 'Limit 2', 'Third attempt', 'Blocked'),
 ('QZ-06', 'Short answer review', M, I, 'An attempt with a short answer', 'Quiz → Attempts → Review → set points', 'Instructor\'s score is final; attempt re-scored'),
 ('QZ-07', 'Short answer with AI off', M, L, 'No AI provider', 'Submit a short answer', 'Marked pending, not zero'),
 ('QZ-08', 'Quiz % in results', M, A, '', 'Results page (staff) and My results (learner)', 'Quizzes % shown in both'),
 ('QZ-09', 'Another batch\'s quiz', M, L, '', 'Open a quiz link from another batch', '"Quiz not available"'),
],
'08 AI Integrations': [
 ('AIC-01', 'Integrations tab', M, A, '', 'Settings → Integrations', 'Cards for OpenAI, Gemini, Azure OpenAI and Groq; master switch; usage'),
 ('AIC-02', 'Save and test a key', M, A, 'A real key (enter it yourself)', 'Paste key, Save, Test', 'Configured, key masked (e.g. sk-…abcd); "Connected"'),
 ('AIC-03', 'Wrong key', M, A, '', 'Test with a wrong key', 'Clear "rejected the API key" message'),
 ('AIC-04', 'Groq: load models', S, A, 'Groq key', 'Load models from Groq; pick one; Test', 'Model list loads; connected'),
 ('AIC-05', 'Capability map', S, A, '', 'Point tutor and evaluate at different providers; save; reload', 'Choices kept'),
 ('AIC-06', 'Budget hard stop', S, A, '', 'Set a tiny monthly budget with hard stop; use AI twice', 'Second call refused; LMS keeps working'),
 ('AIC-07', 'Master switch off', M, A, '', 'Turn AI off', 'Every AI button disappears for everyone; LMS works'),
 ('AIC-08', 'Only admins configure', M, I, '', 'Open Integrations', 'Not available / view only'),
 ('AIC-09', 'Keys never in the browser', M, A, 'Key saved', 'Download a backup; search for the key', 'Not present'),
],
'09 AI Learning Features': [
 ('AI-01', 'Tutor answers from the course', M, L, 'Provider configured', 'Session → SDC Learn AI → ask about the session', 'Answer with Sources; disclaimer shown'),
 ('AI-02', 'Tutor refuses other topics', M, L, '', 'Ask something unrelated', '"Not covered in the approved course materials…"'),
 ('AI-03', 'Rate and escalate', S, L, '', 'Helpful / Escalate', 'Toast; instructor of my batch notified'),
 ('AI-04', 'Tutor keyboard', N, L, '', 'Open drawer with Enter; type; press Enter; press Escape', 'Focus in the box; Enter sends; Escape closes'),
 ('AI-05', 'Summarize a resource', M, L, '', 'Resource → Summarize', 'English summary with AI badge'),
 ('AI-06', 'Pin a summary', S, I, '', 'Summarize → Pin as official summary', 'Learners see "Key concepts" on the session'),
 ('AI-07', 'Practice quiz', M, L, '', 'Test my understanding → generate → answer', 'Score with explanations; not in the gradebook'),
 ('AI-08', 'Quiz generator', M, I, 'Batch workspace', 'Generate with SDC Learn AI → untick one → add', 'Only selected questions; quiz stays draft'),
 ('AI-09', 'Draft evaluator (readable file)', M, I, 'A PDF or text submission', 'Grade → SDC Learn AI → Evaluate', 'Draft feedback and mark filled; nothing saved until Save'),
 ('AI-10', 'Draft evaluator (unreadable file)', M, I, 'An .xlsx submission', 'Evaluate', 'Checklist comments only, no suggested mark'),
 ('AI-11', 'AI-assisted label', S, L, 'Graded with an AI draft', 'My Assignments', '"AI-assisted review" badge'),
 ('AI-12', 'AI unavailable', M, L, 'No provider', 'Ask the tutor', 'Clear "not configured" message; course still works'),
],
'10 AI Course Builder': [
 ('CB-01', 'Off by default', M, A, 'Fresh settings', 'Manage Courses', 'No "Generate with SDC Learn AI" button; direct link shows "not available"'),
 ('CB-02', 'Turn it on', M, A, '', 'Settings → Features → AI course generator on', 'Button appears on Manage Courses'),
 ('CB-03', 'Who can use it', M, I, 'Feature on', 'Check as instructor (no course create) and learner', 'No button / not available'),
 ('CB-04', 'Brief validation', M, A, '', 'Leave topic or audience empty; ask for more than 48 sessions; paste > 12,000 characters', 'Each blocked on the form before any AI request'),
 ('CB-05', 'Generate outline', M, A, 'Provider configured (or demo mode on a developer machine)', 'Fill the brief → Generate', 'Editable outline: title, outcomes, modules, sessions; warnings listed'),
 ('CB-06', 'Edit the outline', M, A, '', 'Add / remove / reorder outcomes, modules, sessions; edit text', 'Warnings recompute; over-limit text blocks saving'),
 ('CB-07', 'Review and create draft', M, A, 'safe saving installed on the database (Settings shows it / ask IT)', 'Review → Create draft course', 'Course saved as Draft with an "AI draft" label; sessions created; nothing published, no fee/instructor/Zoom'),
 ('CB-08', 'Save twice / reload during save', S, A, '', 'Press Create twice quickly; or reload while saving', 'Only one course is created'),
 ('CB-09', 'Discard', S, A, '', 'Discard draft', 'Draft removed; nothing saved'),
 ('CB-10', 'Content tab', M, A, 'AI draft course', 'Course Builder → Content', 'Provider line, completeness "N of M", course check'),
 ('CB-11', 'Generate full content', M, A, 'Provider configured', 'Generate Full Course Content with AI → confirm', 'Parts written one by one and saved after each; progress visible'),
 ('CB-12', 'Pause and switch provider', M, A, 'A provider that hits quota or a wrong key', 'Start generation; when it pauses, Change provider → continue', 'Only missing parts are written; Settings unchanged'),
 ('CB-13', 'Your edits are kept', M, A, 'Content generated', 'Edit a lesson; regenerate that session', 'Your edit is not overwritten'),
 ('CB-14', 'Repair with AI', S, A, '', 'Repair with AI on a field / code check', 'Suggestion shown; applied only if accepted'),
 ('CB-15', 'Code checks', S, A, 'A lesson with code', 'View the session', 'Problems listed; Approve disabled until resolved'),
 ('CB-16', 'Approve a session', S, A, 'All three parts present', 'Approve', 'Approved; new AI text clears approval, your edit keeps it'),
 ('CB-17', 'Create a draft quiz / assignment from content', M, A, 'Course has a batch', 'View → create draft quiz in a batch', 'Draft created in that batch only; not visible to learners until published'),
 ('CB-18', 'Course check with AI', S, A, '', 'Check the course with AI', 'Outcome coverage, missing topics, repetition'),
 ('CB-19', 'Learners see content only when published', M, L, 'Draft course with content', 'Learner opens the course (after enrolment)', 'Nothing visible until the course is published'),
 ('CB-20', 'Storage room check', S, A, 'Large course', 'Generate content near the storage limit', 'Asks first if results may not fit; refuses a save above the limit with a clear message'),
 ('CB-21', 'YouTube key in Integrations', M, A, 'A YouTube Data API v3 key (enter it yourself)', 'Settings → Integrations → YouTube videos: paste key, Save key, Test; then try a wrong key', 'Configured, key masked; Test finds videos; a wrong key gives "YouTube refused the API key"'),
 ('CB-22', 'Videos found automatically', M, A, 'YouTube key saved; AI provider configured', 'Generate an outline → Review', '"Finding a video for each session…", then each session shows a thumbnail, title, channel and length; "N of M sessions" count'),
 ('CB-23', 'No repeated videos', S, A, 'Review step with videos', 'Compare the videos of all sessions', 'No video appears in two sessions of the course'),
 ('CB-24', 'Another / Remove / Paste link', M, A, 'Review step with videos', 'Another video on one session; Remove on another; Paste a YouTube link (try a non-YouTube link first)', 'Video swaps; removed shows "No video"; non-YouTube link refused; pasted video shown'),
 ('CB-25', 'Find videos again', S, A, 'Some sessions without a video', 'Find videos for sessions without one', 'Only empty sessions are searched; others keep their video'),
 ('CB-26', 'Videos saved with the draft', M, A, 'Review with videos', 'Create draft course → Sessions tab → open a session as learner preview', 'Sessions list shows "video"; session page plays the YouTube video with its title and channel under it'),
 ('CB-27', 'No YouTube key', M, A, 'No YouTube key set', 'Generate an outline → Review', 'Clear note that no key is set; draft can still be created without videos'),
 ('CB-28', 'Find videos for an existing course', S, A, 'A course with sessions that have no video', 'Course Builder → Content → Find YouTube videos (N) → confirm', 'Videos attached to empty sessions only; existing videos/recordings unchanged'),
 ('CB-29', 'Changing a found video', S, A, 'A session with a found video', 'Sessions tab → edit the session → change the video URL → save', 'New video plays; the old title/channel line is gone'),
 ('CB-30', 'Who can search videos', M, I, 'Feature on', 'As learner / instructor without course create, try the generator', 'No access; only course creators (and editors in the Content tab) can search'),
],
'11 Sync, Storage & Data': [
 ('SYN-01', 'Status chip', M, A, 'Live site', 'Look at the chip in the top bar after a change', '"Connected — all changes are saved to the server"'),
 ('SYN-02', 'Two devices', M, A, 'Two browsers signed in', 'Change a course title in one; reload the other', 'Change appears'),
 ('SYN-03', 'Offline', S, A, '', 'Go offline, make a change, go online', 'Warning while offline; change saved after reconnecting'),
 ('SYN-04', 'Conflicting edits', S, A, 'Two browsers', 'Edit the same record in both', 'No silent loss; conflict shown'),
 ('SYN-05', 'Session expired', S, L, '', 'Wait for expiry (or clear token) then save', 'Banner asks to sign in again; unsaved changes kept'),
 ('SYN-06', 'Storage meter', S, A, '', 'Fill storage above 80 %', 'Warning bar with Download backup'),
 ('SYN-07', 'Public certificate check', M, G, '', 'verify.html: type "sdc 2026 1"; a revoked code; a random code', 'Auto-formats to SDC-2026-0001; Valid / Revoked / Not found'),
 ('SYN-08', 'Uploads served safely', S, A, '', 'Open an uploaded file link', 'Downloads (not opened as a page)'),
],
'12 Security & Performance': [
 ('SEC-01', 'Data not readable without signing in', M, G, '', 'Call the Supabase table with the public key (IT)', 'Permission denied'),
 ('SEC-02', 'Learner cannot escalate', M, L, '', 'Try to change own role / settings (devtools) and sync', 'Server ignores it'),
 ('SEC-03', 'Only own password hash in browser', M, L, '', 'Inspect stored data', 'Other users have no password hash'),
 ('SEC-04', 'AI keys locked to admins', M, I, '', 'Call /api/ai/configure with an instructor session (IT)', '403'),
 ('SEC-05', 'Site not flagged', M, G, '', 'Open the site in Chrome', 'No "Dangerous site" warning'),
 ('PERF-01', 'Login speed', N, G, 'Empty cache', 'Open login page', 'Ready in about 2 s or less'),
 ('PERF-02', 'Sign-in to My Courses', N, L, '', 'Sign in', 'About 1–2 s'),
 ('PERF-03', 'Page switching', N, A, '', 'Move between pages', 'Instant (well under half a second)'),
],
}

HDR = ['Case ID', 'Area', 'Feature / Workflow', 'Priority', 'Role / Account', 'Preconditions', 'Test steps', 'Expected result', 'Actual result', 'Status', 'Tester', 'Date', 'Evidence / Notes']
thin = Side(style='thin', color='D1D5DB'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL, TITLE = PatternFill('solid', fgColor='0F2A4A'), Font(bold=True, size=14, color='0F2A4A')

wb = Workbook()
ins = wb.active; ins.title = '00 Instructions'
rows = [
 ['SDC Learn — Full Manual Test Suite'], [],
 ['What', 'Every manual test for the whole platform: sign-in, learner, instructor, coordinator, settings, quizzes, all AI features (incl. the AI Course Builder), sync, security and performance.'],
 ['Live site', 'https://ead-university-portal-ten.vercel.app (sign-in: /login.html · verify: /verify.html)'],
 ['Local', 'npm start → http://localhost:3000. Local servers never touch live data. For the AI Course Builder without a key: SDC_AI_MOCK=1 npm start (labelled demo sample, not AI).'],
 ['Demo accounts', 'Password Demo@123 — Coordinator admin@sdclearn.demo · Instructor instructor@sdclearn.demo · Learner learner@sdclearn.demo · Accounts Officer accounts@sdclearn.demo'],
 ['AI tests', 'Need a provider key in Settings → Integrations (enter it yourself). Without one, mark AI cases Blocked.'],
 ['AI Course Builder', 'Off by default: Settings → Features → AI course generator. Saving a draft course needs the database function sdc_state_put_if (supabase/migrations/20261014000100_sdc_state_put_if.sql); /api/sdc/status shows "safeWrite": true when it is installed.'],
 ['Careful', 'Do not run "Reset demo data" or destructive tests on live data without agreement.'],
 ['How to use', 'Work sheet by sheet. Set Status (Pass / Fail / Blocked / N/A / Not Run), write the actual result, attach evidence. The Dashboard counts update automatically.'],
]
for r in rows: ins.append(r)
ins['A1'].font = TITLE
for r in range(3, ins.max_row + 1): ins.cell(r, 1).font = Font(bold=True); ins.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
ins.column_dimensions['A'].width = 20; ins.column_dimensions['B'].width = 120

dash = wb.create_sheet('99 Dashboard')
dash.append(['Execution dashboard']); dash['A1'].font = TITLE
dash.append(['Sheet', 'Total', 'Pass', 'Fail', 'Blocked', 'N/A', 'Not Run', 'Pass %'])
for c in dash[2]: c.font = Font(bold=True, color='FFFFFF'); c.fill = HEAD_FILL

status_dv = '"Pass,Fail,Blocked,N/A,Not Run"'
total = 0
for title, cases in SHEETS.items():
    ws = wb.create_sheet(title)
    ws.append([title]); ws['A1'].font = TITLE
    ws.append(HDR)
    for c in ws[2]: c.font = Font(bold=True, color='FFFFFF'); c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BORDER
    area = title.split(' ', 1)[1]
    for (cid, feat, pri, role, pre, steps, exp) in cases:
        ws.append([cid, area, feat, pri, role, pre, steps, exp, '', 'Not Run', '', '', ''])
    n = len(cases); total += n
    for r in range(3, 3 + n):
        for c in range(1, 14):
            cell = ws.cell(r, c); cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.border = BORDER
    dv = DataValidation(type='list', formula1=status_dv, allow_blank=True); ws.add_data_validation(dv); dv.add(f'J3:J{2 + n}')
    for val, color in (('Pass', 'D1FAE5'), ('Fail', 'FEE2E2'), ('Blocked', 'FFEDD5'), ('N/A', 'E5E7EB')):
        ws.conditional_formatting.add(f'J3:J{2 + n}', CellIsRule(operator='equal', formula=[f'"{val}"'], fill=PatternFill('solid', fgColor=color)))
    for col, w in zip('ABCDEFGHIJKLM', (10, 16, 28, 9, 22, 28, 48, 48, 30, 11, 12, 11, 26)): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'
    q = f"'{title}'!J3:J{2 + n}"
    r = dash.max_row + 1
    dash.append([title, n, f'=COUNTIF({q},"Pass")', f'=COUNTIF({q},"Fail")', f'=COUNTIF({q},"Blocked")', f'=COUNTIF({q},"N/A")', f'=COUNTIF({q},"Not Run")', f'=IFERROR(C{r}/(B{r}-F{r}),0)'])
    dash.cell(r, 8).number_format = '0%'
last = dash.max_row
dash.append(['TOTAL', f'=SUM(B3:B{last})', f'=SUM(C3:C{last})', f'=SUM(D3:D{last})', f'=SUM(E3:E{last})', f'=SUM(F3:F{last})', f'=SUM(G3:G{last})', f'=IFERROR(C{last + 1}/(B{last + 1}-F{last + 1}),0)'])
dash.cell(last + 1, 8).number_format = '0%'
for c in dash[last + 1]: c.font = Font(bold=True)
dash.append([]); dash.append(['Release gate: every Must case Pass (or N/A with a reason). Any Fail blocks release.'])
dash.append(['Tester:']); dash.append(['Date:']); dash.append(['Build / commit:'])
dash.column_dimensions['A'].width = 30
wb.save(OUT)
print(f'{total} cases in {len(SHEETS)} sheets -> {OUT}')
